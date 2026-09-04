"""
Tracegate Report Generation & Finding Selection Acceptance Test Suite
Verifies Tests 1-14 from Section 49 of the Technical Specification:
- Test 1: Project with 1 confirmed Critical finding -> DOCX report generated -> 1 confirmed vulnerability, not 0.
- Test 2: Project with 4 confirmed findings (Critical, High, Low, Info) -> Select only Low and Info -> DOCX contains only Low and Info.
- Test 3: Finding with PoC text -> PoC in DOCX with proper code block formatting (#F8FAFC, Consolas).
- Test 4: Finding with screenshot evidence -> actual image embedded in DOCX.
- Test 5: Finding with user-entered remediation -> user's remediation text appears.
- Test 6: Finding with NO user-entered remediation -> recommended remediation from knowledge base appears.
- Test 7: Project with 2 findings -> Select 1 -> DOCX contains exactly 1 detailed finding section.
- Test 8: 0 findings selected -> attempt generation without clean report flag -> HTTP 400 validation error.
- Test 9: 0 confirmed findings -> attempt generation with clean report flag -> Clean assessment report generated.
- Test 10: Checklist item marked NOT_TESTED -> verify it is NOT counted as clean.
- Test 11: Checklist item marked TESTED_NOT_FOUND -> verify it is counted as verified clean control.
- Test 12: Generate report v1.0, change finding selection, generate v1.1 -> both versions in reports history.
- Test 13: Refresh / reload -> finding selection and persisted PoC/evidence remain available in database.
- Test 14: Cross-project finding ID submitted in selected_finding_ids -> backend rejects with 400 Bad Request.
"""

import unittest
import io
import os
import sys
import uuid
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

import docx
from fastapi.testclient import TestClient
from backend.app import app
from backend import database as db

client = TestClient(app)

# Helper 1x1 base64 PNG
TINY_PNG_B64 = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="

def extract_docx_text(docx_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            tables_text.append(" | ".join(c.text for c in row.cells))
    return "\n".join(paragraphs_text) + "\n" + "\n".join(tables_text)


class TestReportGenerationSuite(unittest.TestCase):

    def setUp(self):
        # Create a unique project for testing
        self.proj_name = f"Test Project {uuid.uuid4().hex[:8]}"
        res = client.post("/api/projects", json={
            "name": self.proj_name,
            "target_url": "https://audit-target.local",
            "environment": "Web Application (Staging)",
            "description": "Automated security assessment test harness"
        })
        self.assertIn(res.status_code, [200, 201])
        self.project = res.json()
        self.project_id = self.project["id"]

    def tearDown(self):
        # Clean up created project
        try:
            client.delete(f"/api/projects/{self.project_id}")
        except Exception:
            pass

    def test_01_one_confirmed_critical_finding_docx_reports_one_not_zero(self):
        """Test 1: Project with 1 confirmed Critical finding -> DOCX report generated -> verify DOCX report shows 1 confirmed vulnerability, not 0."""
        # Add 1 Critical finding
        finding_res = client.post(
            f"/api/projects/{self.project_id}/checklist/item-sql-1/finding",
            json={
                "finding_name": "SQL Injection in Order Search",
                "priority": "CRITICAL",
                "cwe": "CWE-89",
                "description": "Order search parameter directly concatenated into SQL query.",
                "poc_text": "' UNION SELECT username, password FROM users --",
                "test_id": "AUTH-01"
            }
        )
        self.assertIn(finding_res.status_code, [200, 201])
        f_data = finding_res.json()
        finding_id = f_data["id"]

        # Generate report with this finding selected
        rep_res = client.post(
            f"/api/projects/{self.project_id}/reports",
            json={
                "version": "v1.0",
                "author_name": "QA Auditor",
                "selected_finding_ids": [finding_id]
            }
        )
        self.assertIn(rep_res.status_code, [200, 201])
        rep_data = rep_res.json()
        self.assertEqual(rep_data["findings_count"], 1)

        # Download docx
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_data['id']}/download")
        self.assertEqual(dl_res.status_code, 200)
        
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("SQL Injection in Order Search", docx_text)
        self.assertNotIn("0 security vulnerabilities", docx_text)
        self.assertIn("The assessment identified 1 confirmed security finding", docx_text)
        self.assertIn("CRITICAL | 1 | Immediate", docx_text)

    def test_02_select_only_low_and_info_findings(self):
        """Test 2: Project with 4 confirmed findings (Critical, High, Low, Info) -> Select only Low and Info -> Generate DOCX -> Verify report contains only Low and Info."""
        # Add 4 findings of varying severities
        f_crit = client.post(f"/api/projects/{self.project_id}/checklist/item-c/finding", json={
            "finding_name": "Critical Remote Code Execution", "priority": "CRITICAL", "test_id": "AUTH-01"
        }).json()
        f_high = client.post(f"/api/projects/{self.project_id}/checklist/item-h/finding", json={
            "finding_name": "High Privilege Escalation", "priority": "HIGH", "test_id": "AUTH-02"
        }).json()
        f_low = client.post(f"/api/projects/{self.project_id}/checklist/item-l/finding", json={
            "finding_name": "Low Cookie Missing Secure Flag", "priority": "LOW", "test_id": "SESS-01"
        }).json()
        f_info = client.post(f"/api/projects/{self.project_id}/checklist/item-i/finding", json={
            "finding_name": "Info Server Header Disclosure", "priority": "INFORMATIONAL", "test_id": "CONF-01"
        }).json()

        # Select only Low and Info
        rep_res = client.post(
            f"/api/projects/{self.project_id}/reports",
            json={
                "version": "v1.0",
                "author_name": "QA Auditor",
                "selected_finding_ids": [f_low["id"], f_info["id"]]
            }
        )
        self.assertIn(rep_res.status_code, [200, 201])
        rep_data = rep_res.json()
        self.assertEqual(rep_data["findings_count"], 2)
        self.assertEqual(rep_data["info_count"], 1)

        # Verify DOCX content
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_data['id']}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("Low Cookie Missing Secure Flag", docx_text)
        self.assertIn("Info Server Header Disclosure", docx_text)
        self.assertNotIn("Critical Remote Code Execution", docx_text)
        self.assertNotIn("High Privilege Escalation", docx_text)

    def test_03_poc_formatting_in_docx(self):
        """Test 3: Finding with PoC text -> verify PoC appears in DOCX with proper code block formatting (#F8FAFC, Consolas)."""
        poc_payload = "GET /api/v1/auth/reset?user=admin HTTP/1.1\nHost: target.local"
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-poc/finding", json={
            "finding_name": "Password Reset Token Flaw",
            "priority": "HIGH",
            "poc_text": poc_payload,
            "test_id": "AUTH-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        doc = docx.Document(io.BytesIO(dl_res.content))

        # Check for Consolas font and F8FAFC callout shading in table cells
        found_consolas_poc = False
        found_shading = False
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if "GET /api/v1/auth/reset?user=admin" in p.text:
                            for r in p.runs:
                                if r.font.name == "Consolas":
                                    found_consolas_poc = True
                            if "F8FAFC" in cell._tc.xml:
                                found_shading = True

        self.assertTrue(found_consolas_poc, "Consolas-formatted PoC code run should exist in DOCX")
        self.assertTrue(found_shading, "#F8FAFC shaded callout background should exist in PoC block")

    def test_04_screenshot_evidence_embedded_as_image(self):
        """Test 4: Finding with screenshot evidence -> verify image is embedded in DOCX, not just filename."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-evid/finding", json={
            "finding_name": "Reflected XSS on Query Parameter",
            "priority": "MEDIUM",
            "evidence_filename": "xss_alert_box.png",
            "evidence_data": TINY_PNG_B64,
            "test_id": "INP-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        doc = docx.Document(io.BytesIO(dl_res.content))

        # Must have at least 1 inline shape (the embedded image)
        self.assertGreaterEqual(len(doc.inline_shapes), 1, "DOCX must contain embedded image figure")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("Figure 1: Evidence demonstrating the reported behavior", docx_text)

    def test_05_user_remediation_appears(self):
        """Test 5: Finding with user-entered remediation -> verify user's remediation text appears in DOCX."""
        custom_remediation = "Implement strict CSP and HTML entity encoding using DOMPurify."
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-rem/finding", json={
            "finding_name": "Stored XSS in Forum Comments",
            "priority": "HIGH",
            "remediation": custom_remediation,
            "test_id": "INP-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn(custom_remediation, docx_text)

    def test_06_knowledge_base_fallback_remediation(self):
        """Test 6: Finding with NO user-entered remediation -> verify recommended remediation from knowledge base appears."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-norem/finding", json={
            "finding_name": "Unauthenticated API Access",
            "priority": "CRITICAL",
            "cwe": "CWE-306",
            "test_id": "AUTH-01"
            # No remediation provided
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        # Should contain curated remediation from AUTH-01 knowledge base or CWE guidance
        self.assertIn("Technical Remediation", docx_text)
        self.assertIn("Defense-in-Depth Architectural Mitigation", docx_text)

    def test_07_partial_selection_exact_count(self):
        """Test 7: Project with 2 findings -> Select 1 -> verify DOCX contains exactly 1 detailed finding section."""
        f1 = client.post(f"/api/projects/{self.project_id}/checklist/item-p1/finding", json={
            "finding_name": "Active Finding Included in Scope",
            "priority": "HIGH",
            "test_id": "AUTH-01"
        }).json()
        f2 = client.post(f"/api/projects/{self.project_id}/checklist/item-p2/finding", json={
            "finding_name": "Deferred Finding Excluded from Scope",
            "priority": "LOW",
            "test_id": "CONF-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f1["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("Active Finding Included in Scope", docx_text)
        self.assertNotIn("Deferred Finding Excluded from Scope", docx_text)
        self.assertIn("The assessment identified 1 confirmed security finding", docx_text)
        self.assertIn("HIGH | 1 | Urgent", docx_text)

    def test_08_empty_selection_without_clean_flag_rejected(self):
        """Test 8: 0 findings selected -> attempt generation without clean report flag -> verify friendly validation error (HTTP 400)."""
        res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [],
            "allow_clean_report": False
        })
        self.assertEqual(res.status_code, 400)
        data = res.json()
        self.assertIn("at least one", data["detail"].lower())

    def test_09_clean_assessment_report_generated(self):
        """Test 9: 0 confirmed findings -> attempt generation with clean report flag -> verify clean assessment report generated."""
        res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "author_name": "Lead Assessor",
            "selected_finding_ids": [],
            "allow_clean_report": True
        })
        self.assertIn(res.status_code, [200, 201])
        rep_id = res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("No confirmed vulnerabilities were identified", docx_text)
        self.assertIn("Clean Assessment Report (0 vulnerabilities identified)", docx_text)

    def test_10_not_tested_not_counted_as_clean(self):
        """Test 10: Checklist item marked NOT_TESTED -> verify it is NOT counted as clean."""
        # Create a custom test item in project with required fields
        item_res = client.post(
            f"/api/projects/{self.project_id}/checklist/custom",
            json={
                "name": "Brute Force Protection Control",
                "priority": "HIGH",
                "testing_objective": "Verify rate limiting on failed authentication attempts."
            }
        )
        self.assertIn(item_res.status_code, [200, 201])
        item_id = item_res.json()["id"]

        # Update checklist item to NOT_TESTED
        chk_res = client.put(f"/api/checklist/{item_id}/status", json={
            "status": "NOT_TESTED"
        })
        self.assertEqual(chk_res.status_code, 200)

        # Generate report and verify clean controls count
        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "allow_clean_report": True
        })
        self.assertIn(rep_res.status_code, [200, 201])
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_res.json()['id']}/download")
        docx_text = extract_docx_text(dl_res.content)
        # Check that Verified clean controls is 0
        self.assertIn("Verified clean controls: 0", docx_text)

    def test_11_tested_not_found_counted_as_verified_clean(self):
        """Test 11: Checklist item marked TESTED_NOT_FOUND -> verify it is counted as verified clean control."""
        # Create a custom test item in project with required fields
        item_res = client.post(
            f"/api/projects/{self.project_id}/checklist/custom",
            json={
                "name": "MFA Enforcement Control",
                "priority": "HIGH",
                "testing_objective": "Verify multi-factor authentication cannot be bypassed."
            }
        )
        self.assertIn(item_res.status_code, [200, 201])
        item_id = item_res.json()["id"]

        # Update checklist item to TESTED_NOT_FOUND
        chk_res = client.put(f"/api/checklist/{item_id}/status", json={
            "status": "TESTED_NOT_FOUND"
        })
        self.assertEqual(chk_res.status_code, 200)

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "allow_clean_report": True
        })
        self.assertIn(rep_res.status_code, [200, 201])
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_res.json()['id']}/download")
        docx_text = extract_docx_text(dl_res.content)
        # Should report 1 verified clean control and name the control
        self.assertIn("Verified clean controls: 1", docx_text)
        self.assertIn("MFA Enforcement Control", docx_text)

    def test_12_multiple_report_versions_preserved_in_history(self):
        """Test 12: Generate report v1.0, change finding selection, generate v1.1 -> verify both versions are listed in reports history."""
        f1 = client.post(f"/api/projects/{self.project_id}/checklist/item-v1/finding", json={
            "finding_name": "First Finding", "priority": "HIGH", "test_id": "AUTH-01"
        }).json()
        f2 = client.post(f"/api/projects/{self.project_id}/checklist/item-v2/finding", json={
            "finding_name": "Second Finding", "priority": "MEDIUM", "test_id": "AUTH-02"
        }).json()

        # Generate v1.0 with f1
        rep1 = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0", "selected_finding_ids": [f1["id"]]
        }).json()

        # Generate v1.1 with f1 and f2
        rep2 = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.1", "selected_finding_ids": [f1["id"], f2["id"]]
        }).json()

        history_res = client.get(f"/api/projects/{self.project_id}/reports")
        self.assertEqual(history_res.status_code, 200)
        history = history_res.json()
        versions = [r["version"] for r in history]
        self.assertIn("v1.0", versions)
        self.assertIn("v1.1", versions)

    def test_13_persisted_poc_and_evidence_in_database(self):
        """Test 13: Refresh / reload -> finding selection and persisted PoC/evidence remain available in database."""
        poc_text = "curl -X POST https://audit-target.local/api/admin -d 'admin=1'"
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-persist/finding", json={
            "finding_name": "Admin Privilege Escalation",
            "priority": "CRITICAL",
            "poc_text": poc_text,
            "evidence_filename": "admin_escalation.png",
            "evidence_data": TINY_PNG_B64,
            "test_id": "AUTH-01"
        }).json()

        # Fetch findings from backend
        findings_res = client.get(f"/api/projects/{self.project_id}/findings")
        self.assertEqual(findings_res.status_code, 200)
        raw = findings_res.json()
        findings = raw.get("findings", raw) if isinstance(raw, dict) else raw
        match = next((item for item in findings if item["id"] == f["id"]), None)
        self.assertIsNotNone(match)
        self.assertEqual(match["poc_text"], poc_text)
        self.assertEqual(match["evidence_filename"], "admin_escalation.png")
        self.assertEqual(match["evidence_data"], TINY_PNG_B64)

    def test_14_cross_project_finding_id_rejected(self):
        """Test 14: Cross-project finding ID submitted in selected_finding_ids -> backend rejects with 400 Bad Request."""
        # Create Project B
        proj_b_res = client.post("/api/projects", json={
            "name": "Project B", "target_url": "https://target-b.local"
        })
        proj_b_id = proj_b_res.json()["id"]

        # Add finding to Project B
        f_b = client.post(f"/api/projects/{proj_b_id}/checklist/item-b/finding", json={
            "finding_name": "Project B Finding", "priority": "HIGH", "test_id": "AUTH-01"
        }).json()

        try:
            # Try to include Project B's finding into Project A's report
            rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
                "version": "v1.0",
                "selected_finding_ids": [f_b["id"]]
            })
            self.assertEqual(rep_res.status_code, 400)
            data = rep_res.json()
            self.assertIn("not belong to this project", data["detail"])
        finally:
            client.delete(f"/api/projects/{proj_b_id}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
