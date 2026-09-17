"""
backend/certificate_service.py
Tracegate — VAPT Assessment Completion Certificate Service

Automated, professional VAPT Assessment Completion Certificate generation
as the final outcome of the security assessment lifecycle.
Issued strictly when 100% of in-scope findings have passed remediation validation.
"""

import os
import json
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from backend.database import (
    get_project_by_id,
    get_project_findings_list,
    get_db_connection,
    save_certificate_record,
    get_certificate_by_id,
    get_certificate_by_verification_id,
    get_certificates_for_project,
    get_latest_certificate_for_project,
    update_certificate_notice_seen,
    update_certificate_file_paths
)
from backend.report_generator import convert_docx_to_pdf

logger = logging.getLogger("tracegate.certificate")

# Certificate Directories
BASE_DIR = Path(__file__).resolve().parent.parent
CERTIFICATES_DIR = BASE_DIR / "data" / "certificates"
CERTIFICATES_DIR.mkdir(parents=True, exist_ok=True)

# Color Palette (Minimal, formal, professional)
COLOR_NAVY = RGBColor(15, 23, 42)        # Slate/Navy (#0F172A)
COLOR_TEAL = RGBColor(13, 148, 136)      # Deep Teal (#0D9488)
COLOR_SLATE = RGBColor(51, 65, 85)       # Charcoal Slate (#334155)
COLOR_MUTED = RGBColor(100, 116, 139)    # Muted Slate (#64748B)
COLOR_BORDER = RGBColor(226, 232, 240)   # Border Slate (#E2E8F0)
COLOR_CRIT = RGBColor(185, 28, 28)       # Critical (#B91C1C)
COLOR_HIGH = RGBColor(194, 65, 12)       # High (#C2410C)
COLOR_MED = RGBColor(180, 83, 9)         # Medium (#B45309)
COLOR_LOW = RGBColor(4, 120, 87)         # Low (#047857)
COLOR_INFO = RGBColor(2, 132, 199)       # Informational (#0284C7)

HEX_BG_HEADER = "0F172A"
HEX_BG_ACCENT = "F0FDFA"
HEX_BG_SUBTLE = "F8FAFC"
HEX_BG_BORDER = "CBD5E1"

STANDARD_DISCLAIMER = (
    "This certificate confirms completion of the specified assessment and remediation validation "
    "activities recorded by Tracegate. It does not constitute a guarantee of security, immunity from "
    "future vulnerabilities, certification of regulatory compliance, or an independent legal/audit attestation."
)

ATTESTATION_STATEMENT = (
    "This certificate confirms that the security assessment identified below was completed within the "
    "defined assessment scope and that the confirmed findings included in the assessment were successfully "
    "validated through remediation retesting according to the recorded assessment workflow."
)

VALIDATION_SUMMARY_STATEMENT = (
    "All confirmed findings within the defined assessment scope have undergone remediation validation "
    "and were recorded as successfully passed at the time of certificate issuance."
)


# =============================================================================
# XML & STYLING HELPERS
# =============================================================================

def set_cell_background(cell, fill_hex: str):
    """Applies background shading color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top: int = 100, bottom: int = 100, left: int = 140, right: int = 140):
    """Sets internal padding (in dxa) for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def make_table_safe(table):
    """Ensures rows don't split awkwardly across page breaks."""
    for idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cantSplit = parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        trPr.append(cantSplit)
        if idx == 0:
            tblHeader = parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
            trPr.append(tblHeader)


# =============================================================================
# ELIGIBILITY LOGIC
# =============================================================================

def check_assessment_certificate_eligibility(project_id: str) -> Dict[str, Any]:
    """
    Deterministic eligibility check for VAPT Assessment Completion Certificate.
    Eligibility requires:
      1. totalInScopeFindings > 0 (at least one confirmed finding required)
      2. 100% of confirmed findings have status == 'RESOLVED' and retest_status in ('PASSED', 'PASS')
      3. No finding remains in OPEN, IN_REMEDIATION, RETEST_PENDING, or RETEST_FAILED state.
    """
    project = get_project_by_id(project_id)
    if not project:
        return {
            "eligible": False,
            "status": "NOT_FOUND",
            "reason": "Project not found.",
            "total_findings": 0,
            "resolved_findings": 0,
            "pending_retests": 0,
            "failed_retests": 0,
            "open_findings": 0,
            "remaining_by_severity": {},
            "certificate": None
        }

    findings = get_project_findings_list(project_id)
    total_findings = len(findings)

    if total_findings == 0:
        return {
            "eligible": False,
            "status": "NOT_ELIGIBLE",
            "reason": "No confirmed findings recorded for this project. Assessment completion certificate requires at least one in-scope finding.",
            "total_findings": 0,
            "resolved_findings": 0,
            "pending_retests": 0,
            "failed_retests": 0,
            "open_findings": 0,
            "remaining_by_severity": {},
            "certificate": None
        }

    resolved_count = 0
    pending_retests = 0
    failed_retests = 0
    open_findings = 0
    unresolved_findings = []
    remaining_by_severity: Dict[str, int] = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0, "INFORMATIONAL": 0}

    for f in findings:
        f_status = (f.get("status") or "").strip().upper()
        f_retest = (f.get("retest_status") or "").strip().upper()
        sev = (f.get("priority") or "HIGH").strip().upper()
        if sev not in remaining_by_severity:
            sev = "HIGH"

        is_passed = (f_status == "RESOLVED" and "PASS" in f_retest)
        if is_passed:
            resolved_count += 1
        else:
            unresolved_findings.append(f)
            remaining_by_severity[sev] += 1
            if "FAIL" in f_retest:
                failed_retests += 1
            elif "PEND" in f_retest or not f_retest:
                pending_retests += 1
            else:
                open_findings += 1

    # Check if a certificate has already been issued
    existing_cert = get_latest_certificate_for_project(project_id)

    if len(unresolved_findings) == 0:
        status_code = "GENERATED" if existing_cert else "ELIGIBLE"
        return {
            "eligible": True,
            "status": status_code,
            "reason": "All in-scope findings have successfully passed remediation validation.",
            "total_findings": total_findings,
            "resolved_findings": resolved_count,
            "pending_retests": 0,
            "failed_retests": 0,
            "open_findings": 0,
            "remaining_by_severity": remaining_by_severity,
            "certificate": existing_cert
        }
    else:
        unresolved_count = len(unresolved_findings)
        sev_parts = []
        for k in ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFORMATIONAL"]:
            if remaining_by_severity.get(k, 0) > 0:
                sev_parts.append(f"{remaining_by_severity[k]} {k.capitalize()}")
        breakdown = f" ({', '.join(sev_parts)} remaining)" if sev_parts else ""

        if pending_retests > 0 and open_findings == 0 and failed_retests == 0:
            reason = f"{unresolved_count} in-scope finding(s) remain pending retest{breakdown}."
        else:
            reason = f"{unresolved_count} in-scope finding(s) remain unresolved or pending remediation{breakdown}."

        return {
            "eligible": False,
            "status": "NOT_ELIGIBLE",
            "reason": reason,
            "total_findings": total_findings,
            "resolved_findings": resolved_count,
            "pending_retests": pending_retests,
            "failed_retests": failed_retests,
            "open_findings": open_findings,
            "remaining_by_severity": remaining_by_severity,
            "certificate": None
        }


# =============================================================================
# CERTIFICATE GENERATION & PERSISTENCE
# =============================================================================

def generate_assessment_certificate(project_id: str) -> Dict[str, Any]:
    """
    Generate and persist a professional VAPT Assessment Completion Certificate.
    Enforces eligibility, idempotency, and immutability.
    """
    eligibility = check_assessment_certificate_eligibility(project_id)
    if not eligibility["eligible"]:
        return {
            "success": False,
            "status": eligibility["status"],
            "error": eligibility["reason"],
            "eligibility": eligibility
        }

    # Idempotency guard: If already issued for this project and findings state, return existing
    existing = eligibility.get("certificate")
    if existing and existing.get("status") == "VALID":
        snap = existing.get("snapshot") or {}
        if snap.get("total_findings") == eligibility["total_findings"]:
            logger.info(f"Certificate already generated for project {project_id}: {existing['certificate_id']}")
            return {
                "success": True,
                "status": "ALREADY_ISSUED",
                "certificate": existing,
                "eligibility": eligibility
            }

    project = get_project_by_id(project_id)
    findings = get_project_findings_list(project_id)

    # Generate unique IDs
    year = datetime.now().year
    rand_hex = uuid.uuid4().hex[:6].upper()
    cert_id = f"TG-VAPT-{year}-{rand_hex}"

    while get_certificate_by_id(cert_id) is not None:
        rand_hex = uuid.uuid4().hex[:6].upper()
        cert_id = f"TG-VAPT-{year}-{rand_hex}"

    verify_id = f"TG-VERIFY-{uuid.uuid4().hex[:8].upper()}"

    # Determine Associated Assessment / Report ID
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, version, report_title FROM reports WHERE project_id = ? ORDER BY created_at DESC LIMIT 1", (project_id,))
    rep_row = cursor.fetchone()
    conn.close()

    report_id = rep_row["id"] if rep_row else None
    assessment_id = report_id if report_id else project_id

    # Format Dates
    now = datetime.now()
    now_str = now.strftime("%Y-%m-%d %H:%M:%S")
    issue_date = now.strftime("%d %b %Y")

    try:
        p_created = datetime.strptime(project.get("created_at", now_str)[:19], "%Y-%m-%d %H:%M:%S")
        assessment_start = p_created.strftime("%d %b %Y")
    except Exception:
        assessment_start = issue_date

    final_validation_date = issue_date
    assessment_end = issue_date

    # Finding Severity Breakdown
    crit_count = sum(1 for f in findings if (f.get("priority") or "").upper() in ("CRITICAL", "CRIT"))
    high_count = sum(1 for f in findings if (f.get("priority") or "").upper() == "HIGH")
    med_count = sum(1 for f in findings if (f.get("priority") or "").upper() in ("MEDIUM", "MED"))
    low_count = sum(1 for f in findings if (f.get("priority") or "").upper() == "LOW")
    info_count = sum(1 for f in findings if (f.get("priority") or "").upper() in ("INFORMATIONAL", "INFO"))

    # Target & Client
    target_name = project.get("name", "Assessment Target")
    target_url = project.get("target_url") or "Target Web Application"
    client_org = project.get("client_organization") or project.get("organization") or "Not Provided"
    assessment_type = "Web Application Penetration Test (VAPT)"

    # Clean scope summary
    scope_notes = project.get("scope_notes") or project.get("notes") or project.get("description") or "Standard In-Scope Web Application Components"
    clean_scope = scope_notes.strip()[:300]

    # Clean finding snapshot items (NO AI Fix or PR tokens)
    findings_snapshot = []
    for idx, f in enumerate(findings):
        findings_snapshot.append({
            "index": idx + 1,
            "finding_id": f.get("id"),
            "vuln_id": f.get("vuln_id", f"VULN-{idx+1:03d}"),
            "finding_name": f.get("finding_name", "Vulnerability Finding"),
            "severity": (f.get("priority") or "HIGH").upper(),
            "cwe": f.get("cwe", "CWE-200"),
            "status": "RESOLVED",
            "retest_status": "PASSED",
            "retested_at": f.get("updated_at") or now_str
        })

    snapshot_data = {
        "certificate_id": cert_id,
        "verification_id": verify_id,
        "project_id": project_id,
        "assessment_id": assessment_id,
        "report_id": report_id,
        "target_name": target_name,
        "target_url": target_url,
        "client_organization": client_org,
        "assessment_type": assessment_type,
        "assessment_start": assessment_start,
        "assessment_end": assessment_end,
        "issue_date": issue_date,
        "final_validation_date": final_validation_date,
        "total_findings": len(findings),
        "critical_count": crit_count,
        "high_count": high_count,
        "medium_count": med_count,
        "low_count": low_count,
        "info_count": info_count,
        "findings_retested": len(findings),
        "findings_passed": len(findings),
        "findings_failed": 0,
        "findings_summary": findings_snapshot,
        "signatories": {
            "prepared_by": project.get("created_by") or "Security Assessor",
            "validated_by": "Lead Retest Validator",
            "review_status": "Formal Assessment Review Passed"
        },
        "disclaimer": STANDARD_DISCLAIMER
    }

    # Generate Physical Word DOCX and compile to native PDF
    docx_path, pdf_path = build_certificate_documents(snapshot_data)

    # Persist record into vapt_certificates table
    cert_record = {
        "certificate_id": cert_id,
        "verification_id": verify_id,
        "project_id": project_id,
        "assessment_id": assessment_id,
        "report_id": report_id,
        "status": "VALID",
        "issue_date": issue_date,
        "assessment_start": assessment_start,
        "assessment_end": assessment_end,
        "final_validation_date": final_validation_date,
        "total_findings": len(findings),
        "critical_count": crit_count,
        "high_count": high_count,
        "medium_count": med_count,
        "low_count": low_count,
        "info_count": info_count,
        "findings_retested": len(findings),
        "findings_passed": len(findings),
        "findings_failed": 0,
        "target_name": target_name,
        "target_url": target_url,
        "client_organization": client_org,
        "assessment_type": assessment_type,
        "assessment_scope": clean_scope,
        "snapshot": snapshot_data,
        "file_path_docx": str(docx_path) if docx_path else None,
        "file_path_pdf": str(pdf_path) if pdf_path else None,
        "notice_seen": 0,
        "created_at": now_str
    }

    saved = save_certificate_record(cert_record)
    logger.info(f"VAPT Assessment Completion Certificate successfully created: {cert_id}")

    return {
        "success": True,
        "status": "GENERATED",
        "certificate": saved,
        "eligibility": eligibility
    }


# =============================================================================
# PHYSICAL DOCUMENT COMPILATION (DOCX & PDF)
# =============================================================================

def build_certificate_documents(data: Dict[str, Any]) -> Tuple[Optional[Path], Optional[Path]]:
    """
    Constructs a formal, minimalist, high-quality DOCX certificate package
    and compiles native PDF via Microsoft Word COM automation.
    """
    cert_id = data["certificate_id"]
    docx_filename = f"{cert_id}.docx"
    docx_path = CERTIFICATES_DIR / docx_filename

    doc = Document()

    # Document Geometry: Clean portrait with 0.65 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)

    # 1. Formal Border Framing Container (Single cell outer table)
    frame_table = doc.add_table(rows=1, cols=1)
    frame_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer_cell = frame_table.rows[0].cells[0]
    outer_cell.width = Inches(7.2)
    set_cell_background(outer_cell, "FFFFFF")
    set_cell_margins(outer_cell, top=200, bottom=200, left=240, right=240)

    # Apply elegant double border styling to outer frame cell
    tc_pr = outer_cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="double" w:sz="18" w:space="0" w:color="{HEX_BG_HEADER}"/>'
        f'<w:left w:val="double" w:sz="18" w:space="0" w:color="{HEX_BG_HEADER}"/>'
        f'<w:bottom w:val="double" w:sz="18" w:space="0" w:color="{HEX_BG_HEADER}"/>'
        f'<w:right w:val="double" w:sz="18" w:space="0" w:color="{HEX_BG_HEADER}"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    # Content within outer frame
    p_header = outer_cell.paragraphs[0]
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(4)
    p_header.paragraph_format.space_after = Pt(2)

    r_brand = p_header.add_run("TRACEGATE SECURITY WORKSPACE")
    r_brand.font.name = "Calibri"
    r_brand.font.size = Pt(10.5)
    r_brand.font.bold = True
    r_brand.font.color.rgb = COLOR_TEAL

    p_subbrand = outer_cell.add_paragraph()
    p_subbrand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subbrand.paragraph_format.space_before = Pt(0)
    p_subbrand.paragraph_format.space_after = Pt(14)
    r_sub = p_subbrand.add_run("VULNERABILITY ASSESSMENT & PENETRATION TESTING")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(8.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_MUTED

    # Certificate Main Title
    p_title = outer_cell.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("VAPT ASSESSMENT COMPLETION CERTIFICATE")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_subtitle = outer_cell.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_before = Pt(0)
    p_subtitle.paragraph_format.space_after = Pt(12)
    r_st = p_subtitle.add_run("Security Assessment & Remediation Validation Confirmation")
    r_st.font.name = "Calibri"
    r_st.font.size = Pt(10)
    r_st.font.italic = True
    r_st.font.color.rgb = COLOR_TEAL

    # Formal Attestation Statement
    p_attest = outer_cell.add_paragraph()
    p_attest.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_attest.paragraph_format.space_before = Pt(4)
    p_attest.paragraph_format.space_after = Pt(14)
    p_attest.paragraph_format.line_spacing = 1.2
    r_att = p_attest.add_run(ATTESTATION_STATEMENT)
    r_att.font.name = "Calibri"
    r_att.font.size = Pt(9.5)
    r_att.font.color.rgb = COLOR_SLATE

    # Metadata Grid (2-column table)
    meta_table = outer_cell.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_safe(meta_table)

    col_widths = [Inches(3.3), Inches(3.3)]
    for row in meta_table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = col_widths[idx]
            set_cell_background(cell, HEX_BG_SUBTLE)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    rows_data = [
        (("Certificate ID:", data["certificate_id"]), ("Verification ID:", data["verification_id"])),
        (("Assessment ID:", data["assessment_id"]), ("Client / Organization:", data["client_organization"])),
        (("Target Application:", data["target_name"]), ("Target URL / Scope:", data["target_url"])),
        (("Assessment Type:", data["assessment_type"]), ("Assessment Period:", f"{data['assessment_start']} – {data['assessment_end']}")),
        (("Certificate Issue Date:", data["issue_date"]), ("Final Validation Date:", data["final_validation_date"]))
    ]

    for r_idx, ((lbl1, val1), (lbl2, val2)) in enumerate(rows_data):
        c1 = meta_table.rows[r_idx].cells[0]
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r_lbl1 = p1.add_run(f"{lbl1} ")
        r_lbl1.font.bold = True
        r_lbl1.font.size = Pt(8.5)
        r_lbl1.font.color.rgb = COLOR_NAVY
        r_val1 = p1.add_run(str(val1))
        r_val1.font.size = Pt(8.5)
        r_val1.font.color.rgb = COLOR_SLATE

        c2 = meta_table.rows[r_idx].cells[1]
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r_lbl2 = p2.add_run(f"{lbl2} ")
        r_lbl2.font.bold = True
        r_lbl2.font.size = Pt(8.5)
        r_lbl2.font.color.rgb = COLOR_NAVY
        r_val2 = p2.add_run(str(val2))
        r_val2.font.size = Pt(8.5)
        r_val2.font.color.rgb = COLOR_SLATE

    # Spacer
    p_sp1 = outer_cell.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(8)
    p_sp1.paragraph_format.space_after = Pt(4)

    # Findings & Remediation Validation Summary Table
    sum_table = outer_cell.add_table(rows=2, cols=6)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_safe(sum_table)

    headers = ["Total Confirmed", "Critical", "High", "Medium", "Low", "Informational"]
    values = [
        str(data["total_findings"]),
        str(data["critical_count"]),
        str(data["high_count"]),
        str(data["medium_count"]),
        str(data["low_count"]),
        str(data["info_count"])
    ]

    for c_idx, h in enumerate(headers):
        c = sum_table.rows[0].cells[c_idx]
        set_cell_background(c, HEX_BG_HEADER)
        set_cell_margins(c, top=60, bottom=60, left=60, right=60)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for c_idx, v in enumerate(values):
        c = sum_table.rows[1].cells[c_idx]
        set_cell_background(c, HEX_BG_SUBTLE)
        set_cell_margins(c, top=60, bottom=60, left=60, right=60)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(v)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_NAVY

    # Validation Outcome Callout
    p_outcome = outer_cell.add_paragraph()
    p_outcome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_outcome.paragraph_format.space_before = Pt(10)
    p_outcome.paragraph_format.space_after = Pt(10)
    r_out_badge = p_outcome.add_run("✓ REMEDIATION VALIDATION: 100% PASSED (0 REMAINING UNRESOLVED)\n")
    r_out_badge.font.bold = True
    r_out_badge.font.size = Pt(9)
    r_out_badge.font.color.rgb = COLOR_LOW

    r_out_text = p_outcome.add_run(VALIDATION_SUMMARY_STATEMENT)
    r_out_text.font.size = Pt(8.5)
    r_out_text.font.italic = True
    r_out_text.font.color.rgb = COLOR_SLATE

    # Signatories Section (3 columns)
    sign_table = outer_cell.add_table(rows=1, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_safe(sign_table)

    sign_roles = [
        ("PREPARED BY", data["signatories"].get("prepared_by", "Security Assessor"), "Lead Security Assessor"),
        ("VALIDATED BY", data["signatories"].get("validated_by", "Lead Retest Validator"), "VAPT Validation Lead"),
        ("REVIEW STATUS", "Formal Review Completed", data["signatories"].get("review_status", "Assessment Verified"))
    ]

    for c_idx, (role_title, name, subtitle) in enumerate(sign_roles):
        c = sign_table.rows[0].cells[c_idx]
        set_cell_background(c, "FFFFFF")
        set_cell_margins(c, top=80, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        r_rt = p.add_run(f"{role_title}\n")
        r_rt.font.bold = True
        r_rt.font.size = Pt(7.5)
        r_rt.font.color.rgb = COLOR_MUTED

        r_nm = p.add_run(f"{name}\n")
        r_nm.font.bold = True
        r_nm.font.size = Pt(9)
        r_nm.font.color.rgb = COLOR_NAVY

        r_st = p.add_run(subtitle)
        r_st.font.size = Pt(7.5)
        r_st.font.color.rgb = COLOR_MUTED

    # Verification & Authenticity Line
    p_verify = outer_cell.add_paragraph()
    p_verify.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_verify.paragraph_format.space_before = Pt(12)
    p_verify.paragraph_format.space_after = Pt(4)
    r_v = p_verify.add_run(f"Online Verification: /certificate/verify/{data['certificate_id']}   •   Verification Key: {data['verification_id']}")
    r_v.font.size = Pt(8)
    r_v.font.color.rgb = COLOR_TEAL

    # Legal Disclaimer (Section 27)
    p_disc = outer_cell.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disc.paragraph_format.space_before = Pt(4)
    p_disc.paragraph_format.space_after = Pt(2)
    p_disc.paragraph_format.line_spacing = 1.05
    r_disc = p_disc.add_run(data["disclaimer"])
    r_disc.font.size = Pt(7)
    r_disc.font.italic = True
    r_disc.font.color.rgb = COLOR_MUTED

    # Save DOCX
    doc.save(str(docx_path))
    logger.info(f"DOCX certificate generated at: {docx_path}")

    # Compile native PDF via Microsoft Word COM
    pdf_path = None
    try:
        compiled_pdf = convert_docx_to_pdf(str(docx_path))
        if compiled_pdf and Path(compiled_pdf).exists():
            pdf_path = Path(compiled_pdf)
            logger.info(f"PDF certificate generated at: {pdf_path}")
    except Exception as pe:
        logger.warning(f"Could not compile PDF certificate via Word COM: {pe}")

    return docx_path, pdf_path


# =============================================================================
# PUBLIC VERIFICATION SERVICE
# =============================================================================

def get_public_certificate_verification(cert_id_or_verify_id: str) -> Dict[str, Any]:
    """
    Retrieves public verification data for a certificate ID or verification ID.
    STRICT PRIVACY GUARANTEE: Does NOT expose vulnerability names, PoCs, code, tokens,
    or internal private assessment notes.
    """
    cert = get_certificate_by_id(cert_id_or_verify_id)
    if not cert:
        cert = get_certificate_by_verification_id(cert_id_or_verify_id)

    if not cert:
        return {
            "valid": False,
            "status": "NOT_FOUND",
            "message": "Certificate not found in Tracegate authoritative registry."
        }

    return {
        "valid": cert.get("status") == "VALID",
        "status": cert.get("status", "VALID"),
        "certificate_id": cert["certificate_id"],
        "verification_id": cert["verification_id"],
        "target_name": cert["target_name"],
        "target_url": cert["target_url"],
        "client_organization": cert.get("client_organization") or "Not Provided",
        "assessment_type": cert.get("assessment_type", "Web Application Penetration Test (VAPT)"),
        "issue_date": cert["issue_date"],
        "final_validation_date": cert["final_validation_date"],
        "total_findings_validated": cert.get("findings_passed", 0),
        "remediation_validation_status": "100% Passed (Validated)",
        "attestation": ATTESTATION_STATEMENT,
        "disclaimer": cert.get("snapshot", {}).get("disclaimer") or STANDARD_DISCLAIMER
    }
