"""
Tracegate Enterprise VAPT Report & GRC Mapping System Acceptance Test Suite.

Validates the full enterprise report generation lifecycle:
1. Single finding report generation with dual DOCX and native PDF export.
2. Multi-severity distribution across all 5 severity levels (Critical, High, Medium, Low, Informational).
3. Strict finding isolation (verifies zero data leakage/contamination of git branches, PRs, PoCs).
4. Secret redaction scanner (GitHub tokens, Bearer JWTs, passwords, private keys).
5. GRC framework mappings (NIST CSF 2.0, ISO/IEC 27001:2022, SOC 2 TSC, OWASP WSTG v4.2, PCI DSS v4.0).
6. CVSS v4.0 integrity ("CVSS v4.0: Not scored" when vector is unavailable without fabrication).
7. Statutory GRC reference disclaimer verification.
8. Large-scale report performance with 10+ findings and complete appendices.
"""

import unittest
import io
import os
import sys
import uuid
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

import docx
from fastapi.testclient import TestClient
from backend.app import app
from backend.grc_frameworks import MANDATORY_GRC_DISCLAIMER

client = TestClient(app)

TINY_PNG_B64 = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="

def extract_docx_text(docx_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            tables_text.append(" | ".join(c.text for c in row.cells))
    return "\n".join(paragraphs_text) + "\n" + "\n".join(tables_text)


class TestEnterpriseReportSuite(unittest.TestCase):

    def setUp(self):
        self.proj_name = f"Enterprise Audit {uuid.uuid4().hex[:8]}"
        res = client.post("/api/projects", json={
            "name": self.proj_name,
            "target_url": "https://enterprise.tracegate.internal",
            "environment": "Production Banking Gateway",
            "description": "Enterprise security audit and regulatory gap assessment"
        })
        self.assertIn(res.status_code, [200, 201])
        self.project = res.json()
        self.project_id = self.project["id"]

    def tearDown(self):
        try:
            client.delete(f"/api/projects/{self.project_id}")
        except Exception:
            pass

    def test_01_single_finding_report_and_dual_export_docx_and_pdf(self):
        """Verify report generation with 1 Critical finding produces valid DOCX and native Word PDF."""
        finding_res = client.post(
            f"/api/projects/{self.project_id}/checklist/item-sql/finding",
            json={
                "finding_name": "Unauthenticated SQL Injection in Billing API",
                "priority": "CRITICAL",
                "cwe": "CWE-89",
                "description": "SQL injection in billing transaction lookup allows arbitrary database dump.",
                "poc_text": "POST /api/v1/billing HTTP/1.1\nHost: target\n\n{\"id\": \"' OR 1=1--\"}",
                "test_id": "AUTH-01"
            }
        )
        self.assertIn(finding_res.status_code, [200, 201])
        f_id = finding_res.json()["id"]

        rep_res = client.post(
            f"/api/projects/{self.project_id}/reports",
            json={
                "version": "v1.0",
                "author_name": "Lead Pentester",
                "selected_finding_ids": [f_id]
            }
        )
        self.assertIn(rep_res.status_code, [200, 201])
        rep_data = rep_res.json()
        rep_id = rep_data["id"]

        # 1. Download DOCX
        docx_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        self.assertEqual(docx_res.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_res.headers.get("content-type", ""))
        self.assertGreater(len(docx_res.content), 5000)

        docx_text = extract_docx_text(docx_res.content)
        self.assertIn("Unauthenticated SQL Injection in Billing API", docx_text)
        self.assertIn("CRITICAL RISK", docx_text)
        self.assertIn(MANDATORY_GRC_DISCLAIMER, docx_text)

        # 2. Download Native PDF
        pdf_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=pdf")
        self.assertEqual(pdf_res.status_code, 200)
        self.assertIn("application/pdf", pdf_res.headers.get("content-type", ""))
        # Verify PDF header magic bytes %PDF-
        self.assertTrue(pdf_res.content.startswith(b"%PDF-"), "PDF content must start with %PDF- magic bytes")
        self.assertGreater(len(pdf_res.content), 5000)

    def test_02_all_five_severities_reconciliation_and_slas(self):
        """Verify report correctly calculates severity breakdown across all 5 tiers (Crit, High, Med, Low, Info)."""
        severities = [
            ("Remote Code Execution", "CRITICAL", "CWE-89"),
            ("IDOR in Customer Profile", "HIGH", "CWE-639"),
            ("CSRF in Email Update", "MEDIUM", "CWE-352"),
            ("Missing Cache-Control Header", "LOW", "CWE-524"),
            ("Server Version Header Leak", "INFORMATIONAL", "CWE-200")
        ]
        f_ids = []
        for name, prio, cwe in severities:
            f = client.post(f"/api/projects/{self.project_id}/checklist/item-{prio.lower()}/finding", json={
                "finding_name": name,
                "priority": prio,
                "cwe": cwe,
                "test_id": "TEST-01"
            }).json()
            f_ids.append(f["id"])

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": f_ids
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)

        # Check total count and individual counts
        self.assertIn("The assessment identified 5 confirmed security findings requiring remediation.", docx_text)
        self.assertIn("CRITICAL | 1 | Immediate | 24 to 48 Hours | 20.0%", docx_text)
        self.assertIn("HIGH | 1 | Urgent | 7 Calendar Days | 20.0%", docx_text)
        self.assertIn("MEDIUM | 1 | Planned | 30 Calendar Days | 20.0%", docx_text)
        self.assertIn("LOW | 1 | Routine | 90 Calendar Days | 20.0%", docx_text)
        self.assertIn("INFORMATIONAL | 1 | Advisory | Next Scheduled Release | 20.0%", docx_text)

    def test_03_strict_finding_isolation_no_data_contamination(self):
        """Verify Finding A's git branch, PR, and PoC NEVER leak into Finding B's section."""
        f_a = client.post(f"/api/projects/{self.project_id}/checklist/item-a/finding", json={
            "finding_name": "Finding Alpha with AutoFix Branch",
            "priority": "HIGH",
            "poc_text": "EXPLOIT_PAYLOAD_ALPHA_12345",
            "git_branch": "fix/alpha-patch-branch",
            "git_commit": "a1b2c3d4e5",
            "git_pr_url": "https://github.com/org/repo/pull/999",
            "test_id": "AUTH-01"
        }).json()

        f_b = client.post(f"/api/projects/{self.project_id}/checklist/item-b/finding", json={
            "finding_name": "Finding Beta with No Branch",
            "priority": "LOW",
            "poc_text": "EXPLOIT_PAYLOAD_BETA_67890",
            # No git branch, PR, or commit
            "test_id": "CONF-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f_a["id"], f_b["id"]]
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        doc = docx.Document(io.BytesIO(dl_res.content))

        # Inspect paragraphs specifically under Finding Beta
        found_beta = False
        beta_leaked_alpha_git = False
        for p in doc.paragraphs:
            if "Finding Beta with No Branch" in p.text:
                found_beta = True
            if found_beta and ("fix/alpha-patch-branch" in p.text or "pull/999" in p.text or "EXPLOIT_PAYLOAD_ALPHA_12345" in p.text):
                beta_leaked_alpha_git = True
            if "Appendix" in p.text:
                found_beta = False

        self.assertFalse(beta_leaked_alpha_git, "Finding Beta must not contain Finding Alpha's git branch or PoC")

    def test_04_secret_redaction_scanner(self):
        """Verify sensitive credentials (GitHub token, Bearer JWT, database passwords) are redacted."""
        unredacted_poc = (
            "curl -H 'Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIxMjM0NTY3ODkwIn0.DO_NOT_EXPOSE_TOKEN' \\\n"
            "     -H 'X-GitHub-Token: ghp_111122223333444455556666777788889999' \\\n"
            "     https://api.target.com/db_connect?password=SuperSecretPassword123!"
        )

        f = client.post(f"/api/projects/{self.project_id}/checklist/item-sec/finding", json={
            "finding_name": "Hardcoded Secret and Token Exposure",
            "priority": "CRITICAL",
            "cwe": "CWE-798",
            "poc_text": unredacted_poc,
            "test_id": "INFO-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)

        # Raw secrets must NOT be present
        self.assertNotIn("ghp_111122223333444455556666777788889999", docx_text)
        self.assertNotIn("DO_NOT_EXPOSE_TOKEN", docx_text)
        self.assertNotIn("SuperSecretPassword123!", docx_text)

        # Redaction labels must be present
        self.assertIn("[REDACTED GITHUB TOKEN]", docx_text)
        self.assertIn("[REDACTED JWT TOKEN]", docx_text)
        self.assertIn("[REDACTED PASSWORD]", docx_text)

    def test_05_grc_mappings_and_control_gaps(self):
        """Verify deterministic GRC framework mappings and domain gap aggregation."""
        # Add IDOR finding (CWE-639)
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-idor/finding", json={
            "finding_name": "Insecure Direct Object Reference in Invoices",
            "priority": "HIGH",
            "cwe": "CWE-639",
            "test_id": "AUTHZ-04"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)

        # Verify GRC framework mappings for CWE-639
        self.assertIn("PR.AA-05", docx_text)      # NIST CSF 2.0
        self.assertIn("A.8.3", docx_text)         # ISO 27001
        self.assertIn("CC6.1", docx_text)         # SOC 2
        self.assertIn("WSTG-ATHZ-04", docx_text)  # OWASP WSTG
        self.assertIn("7.2", docx_text)           # PCI DSS

        # Verify Observed Control Gaps Domain
        self.assertIn("Broken Object Level Authorization (BOLA / IDOR)", docx_text)

        # Verify Statutory Disclaimer
        self.assertIn(MANDATORY_GRC_DISCLAIMER, docx_text)

    def test_06_cvss_vector_integrity_not_scored(self):
        """Verify CVSS v4.0 explicitly displays 'Not scored' when vector is uncalculated."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-noscore/finding", json={
            "finding_name": "Unvalidated Open Redirect",
            "priority": "LOW",
            "cwe": "CWE-601",
            "test_id": "INP-01"
            # No CVSS score or vector passed
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertIn("CVSS v4.0: Not scored", docx_text)

    def test_07_large_report_ten_plus_findings(self):
        """Verify report engine handles 10+ findings with complete appendices and dual export."""
        f_ids = []
        for i in range(1, 12):
            f = client.post(f"/api/projects/{self.project_id}/checklist/item-bulk-{i}/finding", json={
                "finding_name": f"Comprehensive Bulk Audit Finding #{i}",
                "priority": "HIGH" if i <= 3 else ("MEDIUM" if i <= 7 else "LOW"),
                "cwe": f"CWE-{(80 + i)}",
                "poc_text": f"curl -i https://target/api/v1/test/{i}",
                "test_id": f"BULK-{i}"
            }).json()
            f_ids.append(f["id"])

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v2.0",
            "selected_finding_ids": f_ids
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        # DOCX Verification
        dl_docx = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        self.assertEqual(dl_docx.status_code, 200)
        docx_text = extract_docx_text(dl_docx.content)
        self.assertIn("Comprehensive Bulk Audit Finding #1", docx_text)
        self.assertIn("Comprehensive Bulk Audit Finding #11", docx_text)
        self.assertIn("Appendix A: Complete Findings Register", docx_text)
        self.assertIn("Appendix F: Report Traceability & Audit Metadata", docx_text)

        # PDF Verification
        dl_pdf = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=pdf")
        self.assertEqual(dl_pdf.status_code, 200)
        self.assertTrue(dl_pdf.content.startswith(b"%PDF-"))



    def test_08_coverage_reconciliation_zero_contradiction(self):
        """Verify Problem 1 fix: 12 findings reconcile with coverage table (never shows 0 confirmed vulns or 0% coverage)."""
        f_ids = []
        for i in range(1, 13):
            f = client.post(f"/api/projects/{self.project_id}/checklist/item-rec-{i}/finding", json={
                "finding_name": f"Reconciled Test Finding #{i}",
                "priority": "HIGH" if i <= 4 else ("MEDIUM" if i <= 8 else "LOW"),
                "cwe": "CWE-89" if i % 2 == 0 else "CWE-79",
                "test_id": f"REC-{i}"
            }).json()
            f_ids.append(f["id"])

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": f_ids
        })
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)

        # 12 findings must be reflected in both Executive Summary and Coverage Table
        self.assertIn("The assessment identified 12 confirmed security findings requiring remediation.", docx_text)
        self.assertIn("Controls with confirmed vulnerabilities: 12", docx_text)
        self.assertIn("12 | 12 | 0 | 0 | 12 | 100.0%", docx_text)
        self.assertNotIn("Controls with confirmed vulnerabilities: 0", docx_text)

    def test_09_cvss_formatting_no_contradiction(self):
        """Verify Problem 2 fix: CVSS never outputs 'Score: X (CVSS v4.0: Not scored)'."""
        from backend.report_model import format_cvss_vector

        # 1. Direct unit tests of formatter
        unscored = format_cvss_vector(None, None)
        self.assertEqual(unscored, "CVSS v4.0: Not scored")

        v4_vec = format_cvss_vector(9.8, "CVSS:4.0/AV:N/AC:L/AT:N/PR:N/UI:N/VC:H/VI:H/VA:H/SC:N/SI:N/SA:N")
        self.assertIn("CVSS v4.0 | Score: 9.8", v4_vec)
        self.assertNotIn("Not scored", v4_vec)

        v3_vec = format_cvss_vector(7.5, "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:N/A:N")
        self.assertIn("CVSS v3.1 | Score: 7.5", v3_vec)

        # 2. Docx rendering test
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-cvss/finding", json={
            "finding_name": "CVSS Test Finding",
            "priority": "HIGH",
            "cwe": "CWE-89",
            "test_id": "CVSS-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)
        self.assertNotIn("(CVSS v4.0: Not scored)", docx_text)

    def test_10_remediation_status_normalized_no_fake_pr(self):
        """Verify Problem 3 fix: Finding with 'PR Created' status but no PR url is normalized."""
        from backend.report_model import normalize_remediation_status

        # Without PR metadata, cannot claim 'Pull Request Created'
        norm1 = normalize_remediation_status("PR Created", has_pr=False, has_branch=False, has_commit=False)
        self.assertEqual(norm1, "Fix Proposed (No PR Linked)")

        norm2 = normalize_remediation_status("PR Created", has_pr=False, has_branch=True, has_commit=False)
        self.assertEqual(norm2, "Branch Created (PR Pending)")

        norm3 = normalize_remediation_status("PR Created", has_pr=True, has_branch=True, has_commit=True)
        self.assertEqual(norm3, "Pull Request Created")

    def test_11_tailored_engineering_guidance_sqli_vs_xss(self):
        """Verify Problem 4 fix: SQLi and XSS receive distinct, non-generic technical guidance."""
        from backend.grc_frameworks import map_finding_to_grc_frameworks

        sqli_grc = map_finding_to_grc_frameworks({"cwe": "CWE-89", "finding_name": "SQL Injection"})
        xss_grc = map_finding_to_grc_frameworks({"cwe": "CWE-79", "finding_name": "Reflected XSS"})

        # Root causes must be distinct
        self.assertIn("database query", sqli_grc["technical_root_cause"].lower())
        self.assertIn("document object model", xss_grc["technical_root_cause"].lower())

        # Remediations must be distinct and specific
        self.assertIn("parameterized queries", sqli_grc["remediation"].lower())
        self.assertIn("output encoding", xss_grc["remediation"].lower())

        # SIEM detections must be distinct
        self.assertIn("union select", sqli_grc["siem_detection"].lower())
        self.assertIn("csp", xss_grc["siem_detection"].lower())

        # Developer validations must be distinct
        self.assertIn("sql escape", sqli_grc["developer_validation"].lower())
        self.assertIn("<script>", xss_grc["developer_validation"].lower())

    def test_12_conditional_pci_dss_scoping(self):
        """Verify Problem 5 fix: Non-cardholder projects mark PCI DSS as out-of-scope."""
        # Create non-cardholder project
        hr_proj = client.post("/api/projects", json={
            "name": "Internal Employee HR Portal",
            "target_url": "https://hr.company.internal",
            "environment": "Internal Employee Records (Non-Payment)",
            "description": "HR staff management portal, no cardholder data processed"
        }).json()
        hr_pid = hr_proj["id"]

        try:
            f = client.post(f"/api/projects/{hr_pid}/checklist/item-hr/finding", json={
                "finding_name": "Session Timeout Missing",
                "priority": "LOW",
                "cwe": "CWE-384",
                "test_id": "SESS-01"
            }).json()

            rep_res = client.post(f"/api/projects/{hr_pid}/reports", json={
                "version": "v1.0",
                "selected_finding_ids": [f["id"]]
            })
            rep_id = rep_res.json()["id"]

            dl_res = client.get(f"/api/projects/{hr_pid}/reports/{rep_id}/download")
            docx_text = extract_docx_text(dl_res.content)

            # PCI DSS in non-cardholder project should indicate out-of-scope / N/A
            self.assertIn("Out of Cardholder Data Environment Scope", docx_text)
        finally:
            client.delete(f"/api/projects/{hr_pid}")

    def test_13_sla_language_recommended_target(self):
        """Verify Problem 6 fix: SLA text frames timeframes as recommended priority targets."""
        # Generate report with 1 finding
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-sla/finding", json={
            "finding_name": "SLA Language Test Finding",
            "priority": "HIGH",
            "cwe": "CWE-639",
            "test_id": "SLA-01"
        }).json()
        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download")
        docx_text = extract_docx_text(dl_res.content)

        # Assert non-contractual guidance language
        self.assertIn("Organizations should remediate identified findings according to recommended priority targets", docx_text)
        self.assertIn("They do not constitute contractual or regulatory service level agreements", docx_text)

    def test_14_pre_render_integrity_validator(self):
        """Verify Problem 7 fix: Pre-render validator detects count mismatches and contradictions."""
        from backend.report_model import validate_report_model_integrity

        # Valid model fixture
        valid_model = {
            "metrics": {
                "total_findings": 1,
                "critical_count": 1,
                "high_count": 0,
                "medium_count": 0,
                "low_count": 0,
                "info_count": 0,
                "total_tests": 5,
                "tested_tests": 1,
                "untested_tests": 4,
                "clean_tests": 0,
                "vulnerable_tests": 1,
            },
            "findings": [
                {
                    "report_vuln_id": "VULN-001",
                    "cvss_formatted": "CVSS v4.0 | Score: 9.8 | Vector: CVSS:4.0/AV:N/AC:L",
                    "fix_status": "Remediated"
                }
            ]
        }
        errs = validate_report_model_integrity(valid_model)
        self.assertEqual(len(errs), 0)

        # Invalid model fixture: finding count mismatch & contradictory CVSS string
        invalid_model = {
            "metrics": {
                "total_findings": 2,  # Claims 2, but list has 1
                "critical_count": 1,
                "high_count": 0,
                "medium_count": 0,
                "low_count": 0,
                "info_count": 0,
                "total_tests": 5,
                "tested_tests": 1,
                "untested_tests": 4,
                "clean_tests": 0,
                "vulnerable_tests": 0,  # Contradicts total_findings
            },
            "findings": [
                {
                    "report_vuln_id": "VULN-001",
                    "cvss_formatted": "Score: 9.8 (CVSS v4.0: Not scored)",  # Contradiction
                    "fix_status": "Pull Request Created",  # No PR URL or number
                    "git_pr_url": None,
                    "git_pr_number": None
                }
            ]
        }
        errs = validate_report_model_integrity(invalid_model)
        self.assertGreater(len(errs), 0)
        err_str = " ".join(errs)
        self.assertIn("Finding count mismatch", err_str)
        self.assertIn("Coverage contradiction", err_str)
        self.assertIn("contradictory CVSS string", err_str)
        self.assertIn("claims 'Pull Request Created'", err_str)

    def test_15_rule_80_standalone_separation_no_autofix_leakage(self):
        """Verify Rule 80: Report is completely independent of AI AutoFix; zero AI/git branches leak into client reports."""
        from backend.database import get_db_connection
        
        # 1. Create a finding with AutoFix / Git fields populated in the database
        finding_res = client.post(
            f"/api/projects/{self.project_id}/checklist/item-autofix/finding",
            json={
                "finding_name": "SQL Injection in Search API",
                "priority": "HIGH",
                "cwe": "CWE-89",
                "description": "Unescaped user parameter in search query.",
                "test_id": "AUTH-02"
            }
        )
        self.assertIn(finding_res.status_code, [200, 201])
        f_id = finding_res.json()["id"]

        # Directly update DB record to simulate an autofix run with internal branch & PR
        conn = get_db_connection()
        try:
            conn.execute(
                """
                UPDATE findings
                SET github_branch = ?, github_commit = ?, github_pr = ?, fix_status = ?
                WHERE id = ?
                """,
                ("tracegate/fix/sql-injection-search", "deadbeef12345678", "https://github.com/client-org/repo/pull/99", "Pull Request Created", f_id)
            )
            conn.commit()
        finally:
            conn.close()

        # 2. Generate report
        rep_res = client.post(
            f"/api/projects/{self.project_id}/reports",
            json={
                "version": "v1.0",
                "selected_finding_ids": [f_id]
            }
        )
        self.assertIn(rep_res.status_code, [200, 201])
        rep_id = rep_res.json()["id"]

        # 3. Download DOCX and extract text
        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        self.assertEqual(dl_res.status_code, 200)
        docx_text = extract_docx_text(dl_res.content)

        # Assert mandatory professional client report remediation sections are present
        self.assertIn("Recommended Technical Remediation", docx_text)
        self.assertIn("Developer Validation Guidance", docx_text)

        # Assert zero leakage of internal AI AutoFix features or git branch names
        self.assertNotIn("AI AutoFix", docx_text)
        self.assertNotIn("AI Fix", docx_text)
        self.assertNotIn("tracegate/fix/", docx_text)
        self.assertNotIn("deadbeef12345678", docx_text)
        self.assertNotIn("pull/99", docx_text)

    def test_16_rule_70_finding_id_collision_resolution(self):
        """Verify Rule 70: Finding ID collisions are resolved to unique deterministic IDs (e.g. VULN-001, VULN-002, VULN-003)."""
        from backend.report_model import build_report_model
        
        project = {
            "id": self.project_id,
            "name": self.proj_name,
            "target_url": "https://enterprise.tracegate.internal"
        }
        # Two findings intentionally given colliding vuln_id in raw representation
        raw_findings = [
            {"id": "f-coll-1", "vuln_id": "VULN-002", "finding_name": "Stored XSS", "priority": "HIGH", "cwe": "CWE-79"},
            {"id": "f-coll-2", "vuln_id": "VULN-002", "finding_name": "Reflected XSS", "priority": "HIGH", "cwe": "CWE-79"},
            {"id": "f-coll-3", "vuln_id": "VULN-003", "finding_name": "DOM XSS", "priority": "MEDIUM", "cwe": "CWE-79"}
        ]
        checklist = [
            {"id": "c-1", "status": "VULNERABLE", "finding": raw_findings[0]},
            {"id": "c-2", "status": "VULNERABLE", "finding": raw_findings[1]},
            {"id": "c-3", "status": "VULNERABLE", "finding": raw_findings[2]}
        ]
        model = build_report_model(
            project=project,
            selected_finding_ids=["f-coll-1", "f-coll-2", "f-coll-3"],
            findings=raw_findings,
            checklist_items=checklist
        )
        report_ids = [f["report_vuln_id"] for f in model["findings"]]
        # All report IDs must be strictly unique
        self.assertEqual(len(report_ids), len(set(report_ids)), "All report_vuln_id values must be unique")
        self.assertEqual(len(report_ids), 3)
        self.assertEqual(report_ids[0], "VULN-002")
        self.assertNotEqual(report_ids[1], "VULN-002")
        self.assertTrue(report_ids[1].startswith("VULN-"))

    def test_17_rule_27_structured_developer_remediation(self):
        """Verify Rule 27: Structured developer remediation contains all 5 required subsections."""
        from backend.report_model import build_report_model
        
        project = {
            "id": self.project_id,
            "name": self.proj_name,
            "target_url": "https://enterprise.tracegate.internal"
        }
        raw_findings = [
            {
                "id": "f-struct-1",
                "vuln_id": "VULN-001",
                "finding_name": "SQL Injection Flaw",
                "priority": "CRITICAL",
                "cwe": "CWE-89",
                "description": "User input concatenated directly into SQL statement."
            }
        ]
        checklist = [{"id": "c-1", "status": "VULNERABLE", "finding": raw_findings[0]}]
        model = build_report_model(
            project=project,
            selected_finding_ids=["f-struct-1"],
            findings=raw_findings,
            checklist_items=checklist
        )
        finding = model["findings"][0]
        remediation_text = finding.get("remediation", "")
        
        # Verify the 5 required subsections in structured remediation
        self.assertIn("Root Problem:", remediation_text)
        self.assertIn("Required Code Change:", remediation_text)
        self.assertIn("Security Control:", remediation_text)
        self.assertIn("Implementation Guidance:", remediation_text)
        self.assertIn("Regression Testing:", remediation_text)


    def test_18_authoritative_27_sections_single_render_guarantee(self):
        """Verify Rule 62: All 27 authoritative sections and Appendices A-G render exactly once without duplicates."""
        # 1. Create finding and generate report
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-27sec/finding", json={
            "finding_name": "Stored Cross-Site Scripting in Comment Field",
            "priority": "HIGH",
            "cwe": "CWE-79",
            "test_id": "INPV-02"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        doc = docx.Document(io.BytesIO(dl_res.content))

        # Collect all Heading 1 paragraphs
        h1_texts = [p.text.strip() for p in doc.paragraphs if p.style and "Heading 1" in p.style.name]

        # 27 Sections Verification (Sections 2 through 27)
        expected_sections = [
            "Section 2: Document Control & Revision History",
            "Section 3: Table of Contents & Navigation Map",
            "Section 4: Confidentiality, Legal Disclaimer & Rules of Engagement",
            "Section 5: Executive Summary",
            "Section 6: Assessment Scope & Boundary Definition",
            "Section 7: Target Environment & Architecture Overview",
            "Section 8: Assessment Methodology & Testing Standard",
            "Section 9: Overall Security Posture & Executive Risk Rating",
            "Section 10: Vulnerability Severity Distribution & Metrics",
            "Section 11: Attack Surface & Component Exposure Analysis",
            "Section 12: Testing Coverage & Verification Matrix",
            "Section 13: Summary of Findings (Consolidated Table)",
            "Section 14: Category-Wise Findings Analysis",
            "Section 15: Detailed Technical Vulnerability Findings",
            "Section 16: Retest Status & Verification Lifecycle",
            "Section 17: Vulnerability Remediation Register",
            "Section 18: Prioritized Remediation Action Plan",
            "Section 19: Positive Security Controls & Defenses Observed",
            "Section 20: Defense-in-Depth & Architectural Hardening Recommendations",
            "Section 21: GRC Framework Compliance Mapping (NIST CSF 2.0, ISO/IEC 27001:2022, SOC 2, PCI DSS 4.0)",
            "Section 22: Strategic & Tactical Recommendations for Leadership",
            "Section 23: Risk Acceptance & Residual Risk Guidance",
            "Section 24: Secure Development Lifecycle (SDLC) Integration Recommendations",
            "Section 25: Evidence Register & Screenshot Manifest",
            "Section 26: Assessment Sign-Off & Attestation",
            "Section 27: Appendices"
        ]

        for sec in expected_sections:
            count = h1_texts.count(sec)
            self.assertEqual(count, 1, f"Section '{sec}' must appear exactly once in docx, but appeared {count} times")

        # Appendices A through G Verification
        expected_appendices = [
            "Appendix A: Complete Findings Register & Severity Rating Methodology",
            "Appendix B: Evidence Index & Figure Register",
            "Appendix C: CWE Classification Index",
            "Appendix D: GRC Framework References & Methodology Standards",
            "Appendix E: Retest & Remediation Verification Summary",
            "Appendix F: Report Traceability & Audit Metadata",
            "Appendix G: Glossary of Cybersecurity Terminology"
        ]

        for app in expected_appendices:
            count = h1_texts.count(app)
            self.assertEqual(count, 1, f"Appendix '{app}' must appear exactly once in docx, but appeared {count} times")

    def test_19_report_generation_idempotency(self):
        """Verify report generation is strictly idempotent: 1x, 2x, 3x generation produces consistent output without accumulation."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-idem/finding", json={
            "finding_name": "Idempotency Test Finding",
            "priority": "MEDIUM",
            "cwe": "CWE-352",
            "test_id": "CSRF-01"
        }).json()

        # Generate 3 successive reports
        doc_lengths = []
        for v in ["v1.0", "v1.1", "v1.2"]:
            rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
                "version": v,
                "selected_finding_ids": [f["id"]]
            })
            rep_id = rep_res.json()["id"]
            dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
            docx_text = extract_docx_text(dl_res.content)
            # Finding appears across Consolidated Table, Detailed Finding H2, Finding Metadata, Retest Register, and Evidence Manifest
            self.assertEqual(docx_text.count("Idempotency Test Finding"), 5)
            doc_lengths.append(len(docx_text))
        # Document length remains stable and finding count never accumulates across successive generations
        self.assertEqual(len(doc_lengths), 3)
        self.assertGreater(doc_lengths[0], 10000)
        self.assertLess(abs(doc_lengths[2] - doc_lengths[0]), 500)

    def test_20_rules_of_engagement_and_residual_risk(self):
        """Verify Section 4 Rules of Engagement and Section 23 Residual Risk guidance are rendered."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-roe/finding", json={
            "finding_name": "API Rate Limit Missing",
            "priority": "LOW",
            "cwe": "CWE-307",
            "test_id": "RATE-01"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        docx_text = extract_docx_text(dl_res.content)

        # Section 4 Rules of Engagement assertions
        self.assertIn("Authorized Target Scope", docx_text)
        self.assertIn("Permitted Testing Techniques", docx_text)
        self.assertIn("Prohibited Attack Vectors", docx_text)
        self.assertIn("Data Handling & Encryption", docx_text)

        # Section 23 Residual Risk assertions
        self.assertIn("Current Assessed Residual Risk", docx_text)
        self.assertIn("RISK ACCEPTANCE POLICY", docx_text)

    def test_21_appendix_g_glossary_and_appendices(self):
        """Verify Appendix G Glossary of Cybersecurity Terminology is fully rendered with key terms."""
        f = client.post(f"/api/projects/{self.project_id}/checklist/item-gloss/finding", json={
            "finding_name": "Glossary Verification Finding",
            "priority": "LOW",
            "cwe": "CWE-200",
            "test_id": "INFO-02"
        }).json()

        rep_res = client.post(f"/api/projects/{self.project_id}/reports", json={
            "version": "v1.0",
            "selected_finding_ids": [f["id"]]
        })
        rep_id = rep_res.json()["id"]

        dl_res = client.get(f"/api/projects/{self.project_id}/reports/{rep_id}/download?format=docx")
        docx_text = extract_docx_text(dl_res.content)

        self.assertIn("Appendix G: Glossary of Cybersecurity Terminology", docx_text)
        self.assertIn("VAPT", docx_text)
        self.assertIn("CWE", docx_text)
        self.assertIn("CVSS", docx_text)
        self.assertIn("BOLA / IDOR", docx_text)
        self.assertIn("RCE", docx_text)

    def test_22_registry_blocks_duplicate_section_render(self):
        """Verify ReportSectionRegistry enforces single render guarantee and rejects duplicate calls."""
        from backend.report_generator import ReportSectionRegistry
        from backend.report_model import ReportIntegrityError

        registry = ReportSectionRegistry()
        executed = []

        def sample_render(name):
            executed.append(name)
            return f"rendered_{name}"

        # First call succeeds
        res = registry.render_section("section_test", sample_render, "test")
        self.assertEqual(res, "rendered_test")
        self.assertEqual(executed, ["test"])

        # Duplicate call raises ReportIntegrityError with DUPLICATE_SECTION
        with self.assertRaises(ReportIntegrityError) as ctx:
            registry.render_section("section_test", sample_render, "test")

        self.assertIn("DUPLICATE_SECTION", str(ctx.exception))
        self.assertIn("section_test", str(ctx.exception))
        # Execution function was NOT invoked a second time
        self.assertEqual(len(executed), 1)

    def test_23_verify_five_quality_pass_issues_resolved(self):
        """Verify:
        1. No contradictory CVSS labels ('CVSS v3.1 Scoring & Vector' over 'CVSS v4.0: Not scored').
        2. No double numbering in reproduction steps (zero '1. 1.', '2. 2.').
        3. Endpoint vs source path separation (profileController.py / sqlinjection.py under source path, not bare endpoint).
        4. Stored SVG XSS (VULN-009) observation accurately describes SVG upload & XML execution.
        5. Methodology claim softened to 'conducted with reference to' rather than 'executed strictly in accordance with'.
        """
        mock_proj = {
            "id": "proj-verify-5-issues",
            "name": "Audit Verification Target",
            "target_url": "https://staging.target.local",
            "organization": "Target Security Org",
            "findings": [
                {
                    "id": "f-idor",
                    "vuln_id": "VULN-001",
                    "finding_name": "Insecure Direct Object Reference (IDOR)",
                    "priority": "HIGH",
                    "cwe": "CWE-639",
                    "cvss": None,
                    "cvss_vector": None,
                    "endpoint": "backend/controllers/profileController.py",
                    "reproduction_steps": "1. Navigate to user settings.\\n2. Modify user_id parameter to 205.\\n3. Observe unauthorized profile data.",
                    "poc_text": "GET /api/user/profile?id=205 HTTP/1.1\\nHost: staging.target.local"
                },
                {
                    "id": "f-svg",
                    "vuln_id": "VULN-009",
                    "finding_name": "[VULN-009] Stored XSS via Malicious SVG Image Upload",
                    "priority": "HIGH",
                    "cwe": "CWE-79",
                    "cvss": None,
                    "cvss_vector": None,
                    "affected_component": "backend/controllers/svgController.py",
                    "reproduction_steps": ["1. Upload malicious avatar.svg containing <script>alert(1)</script>", "2. Navigate to /uploads/avatar.svg", "3. Trigger script execution"],
                    "poc_text": "POST /api/upload HTTP/1.1\\nContent-Type: multipart/form-data\\n\\n<svg onload=alert(1)>"
                },
                {
                    "id": "f-sqli",
                    "vuln_id": "VULN-003",
                    "finding_name": "SQL Injection in Authentication",
                    "priority": "CRITICAL",
                    "cwe": "CWE-89",
                    "cvss": 9.8,
                    "cvss_vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
                    "affected_component": "sqlinjection.py",
                    "reproduction_steps": "1. Intercept POST /login\\n2. Inject admin' OR '1'='1 into password\\n3. Authenticate as admin",
                    "poc_text": "POST /api/login HTTP/1.1\\nHost: staging.target.local\\n\\nusername=admin&password=admin' OR '1'='1"
                }
            ],
            "checklist": [
                {"test_id": "T1", "name": "Auth Test", "status": "VULNERABILITY_FOUND"},
                {"test_id": "T2", "name": "Upload Test", "status": "VULNERABILITY_FOUND"},
                {"test_id": "T3", "name": "IDOR Test", "status": "VULNERABILITY_FOUND"},
                {"test_id": "T4", "name": "CSRF Test", "status": "TESTED_NOT_FOUND"},
            ]
        }

        from backend.report_generator import generate_docx_report
        import docx

        docx_path = generate_docx_report(project=mock_proj, version="v1.0", author_name="Lead Tester")
        self.assertTrue(Path(docx_path).exists())

        doc = docx.Document(docx_path)
        all_text = "\\n".join([p.text for p in doc.paragraphs])

        # 1. CVSS Contradiction check:
        for tbl in doc.tables:
            for row in tbl.rows:
                if len(row.cells) == 2:
                    c0 = row.cells[0].text.strip()
                    c1 = row.cells[1].text.strip()
                    if "CVSS v3.1" in c0:
                        self.assertNotIn("CVSS v4.0: Not scored", c1)
                    if "CVSS v4.0: Not scored" in c1:
                        self.assertIn("v4.0", c0)

        # 2. Reproduction numbering check:
        self.assertNotIn("1. 1.", all_text)
        self.assertNotIn("2. 2.", all_text)
        self.assertNotIn("3. 3.", all_text)

        # 3. Asset / Endpoint separation check:
        self.assertNotIn("Target Endpoint / URI: backend/controllers/profileController.py", all_text)
        self.assertIn("Target Endpoint / URI: /api/user/profile", all_text)
        self.assertIn("Source Component / Code Path: backend/controllers/profileController.py", all_text)
        self.assertNotIn("Target Endpoint / URI: sqlinjection.py", all_text)
        self.assertIn("Target Endpoint / URI: /api/login", all_text)

        # 4. Technical Observation for SVG XSS:
        self.assertIn("SVG (Scalable Vector Graphics)", all_text)
        self.assertIn("image/svg+xml", all_text)

        # 5. Methodology claim check:
        self.assertNotIn("executed strictly in accordance with", all_text)
        self.assertIn("conducted with reference to recognized industry standards", all_text)

if __name__ == "__main__":
    unittest.main(verbosity=2)

